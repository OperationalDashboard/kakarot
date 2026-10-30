"""Knows your files: indexes the documents in your allowed folders so Kakarot can answer questions from them.
Uses meaning search (a small local embedding model) plus keyword search, combined."""
import logging
import os
import re
import struct
import threading
import time

import requests

import db
from config import BASE, CFG
from tools import I, ROOTS, S, _cut, tool

log = logging.getLogger("kakarot.knowledge")
_index_lock = threading.Lock()
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".html", ".htm", ".xml", ".py", ".js", ".ini", ".yaml", ".yml"}
DOC_EXT = {".pdf", ".docx", ".xlsx", ".xlsm"}
SKIP_DIRS = {"_trash", "node_modules", ".git", "venv", "__pycache__", "screenshots", "uploads_tmp"}
MAX_FILE_MB = 25
MAX_FILES = 4000
CHUNK, OVERLAP, MAX_CHUNKS = 1200, 200, 300
STATE = {"running": False, "last": None, "last_note": ""}
_cache = {}


def embed_model():
    return CFG.get("embed_model", "nomic-embed-text")


# ---------------------------------------------------------------- reading files
def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(path).pages[:300])
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows)
        return text
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets[:10]:
            out.append(f"Sheet: {ws.title}")
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i >= 2000:
                    break
                out.append(" | ".join("" if v is None else str(v) for v in r))
        return "\n".join(out)
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read(600_000)


def chunks(text):
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    out, i = [], 0
    while i < len(text) and len(out) < MAX_CHUNKS:
        piece = text[i:i + CHUNK]
        if i + CHUNK < len(text):  # end on a line or sentence break when possible
            cut = max(piece.rfind("\n"), piece.rfind(". "))
            if cut > CHUNK * 0.6:
                piece = piece[:cut + 1]
        if piece.strip():
            out.append(piece.strip())
        if i + len(piece) >= len(text):
            break
        i += max(len(piece) - OVERLAP, CHUNK // 2)
    return out


# ---------------------------------------------------------------- embeddings
def embed(texts, kind="document"):
    """Returns a list of vectors, or None if the embedding model isn't available."""
    prefix = "search_query: " if kind == "query" else "search_document: "
    try:
        r = requests.post(CFG["ollama_url"].rstrip("/") + "/api/embed", timeout=300,
                          json={"model": embed_model(), "input": [prefix + t[:2000] for t in texts],
                                "keep_alive": "10m"})
        if r.status_code != 200:
            return None
        return r.json().get("embeddings")
    except Exception:
        return None


def _pack(v):
    return struct.pack(f"{len(v)}f", *v)


def _unpack(b):
    return struct.unpack(f"{len(b) // 4}f", b)


# ---------------------------------------------------------------- indexing
def _walk():
    count = 0
    for root in ROOTS:
        for dirpath, dirs, files in os.walk(root):
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS and not d.startswith(".")
                       and os.path.normcase(os.path.join(dirpath, d)) != os.path.normcase(BASE)]  # not Kakarot itself
            for name in files:
                ext = os.path.splitext(name)[1].lower()
                if ext not in TEXT_EXT and ext not in DOC_EXT or name.startswith(("~$", ".")):
                    continue
                p = os.path.join(dirpath, name)
                try:
                    st = os.stat(p)
                except OSError:
                    continue
                if st.st_size > MAX_FILE_MB * 1024 * 1024:
                    continue
                yield p, st.st_mtime, st.st_size
                count += 1
                if count >= MAX_FILES:
                    return


def _remove(c, path):
    c.execute("DELETE FROM doc_chunks WHERE path=?", (path,))
    try:
        c.execute("DELETE FROM doc_fts WHERE path=?", (path,))
    except Exception:
        pass
    c.execute("DELETE FROM doc_files WHERE path=?", (path,))


def index_files(force=False):
    """Index new/changed files, drop deleted ones. Safe to call often (only changed files are re-read)."""
    if not _index_lock.acquire(blocking=False):
        return "Indexing is already running."
    STATE["running"] = True
    t0 = time.time()
    try:
        known = {r["path"]: r for r in db.q("SELECT path, mtime, size FROM doc_files")}
        seen, new, failed, use_emb = set(), 0, 0, None
        for p, mtime, size in _walk():
            seen.add(p)
            k = known.get(p)
            if k and not force and abs(k["mtime"] - mtime) < 1 and k["size"] == size:
                continue
            try:
                parts = chunks(extract_text(p))
            except Exception as e:
                log.info("skip %s: %s", p, e)
                failed += 1
                parts = []
            vecs = None
            if parts and use_emb is not False:
                vecs = []
                for i in range(0, len(parts), 16):
                    v = embed(parts[i:i + 16])
                    if v is None:
                        vecs, use_emb = None, False
                        break
                    vecs += v
                else:
                    use_emb = True
            with db.tx() as c:
                _remove(c, p)
                for n, text in enumerate(parts):
                    cid = c.execute("INSERT INTO doc_chunks(path, chunk_no, text, emb) VALUES(?,?,?,?)",
                                    (p, n, text, _pack(vecs[n]) if vecs else None)).lastrowid
                    try:
                        c.execute("INSERT INTO doc_fts(text, path, chunk_id) VALUES(?,?,?)", (text, p, cid))
                    except Exception:
                        pass
                c.execute("INSERT OR REPLACE INTO doc_files(path, mtime, size, chunks, indexed) VALUES(?,?,?,?,?)",
                          (p, mtime, size, len(parts), db.now()))
            new += 1
        gone = [p for p in known if p not in seen]
        with db.tx() as c:
            for p in gone:
                _remove(c, p)
        if use_emb is not False:
            use_emb = _backfill()
        note = (f"Indexed {new} new/changed file(s), removed {len(gone)}, skipped {failed} unreadable, "
                f"in {time.time() - t0:.0f}s.")
        if use_emb is False:
            note += f" Meaning search is off until you run get-advanced-models.bat ({embed_model()}); keyword search works."
        STATE.update(last=db.now(), last_note=note)
        return note
    finally:
        STATE["running"] = False
        _index_lock.release()


def _backfill():
    """Give meaning-vectors to passages indexed before the embedding model was downloaded."""
    rows = db.q("SELECT id, text FROM doc_chunks WHERE emb IS NULL LIMIT 5000")
    if not rows:
        return True
    for i in range(0, len(rows), 16):
        part = rows[i:i + 16]
        v = embed([r["text"] for r in part])
        if v is None:
            return False
        with db.tx() as c:
            for r, vec in zip(part, v):
                c.execute("UPDATE doc_chunks SET emb=? WHERE id=?", (_pack(vec), r["id"]))
    return True


def stats():
    r = db.q("SELECT COUNT(*) AS files, COALESCE(SUM(chunks),0) AS chunks, MAX(indexed) AS last FROM doc_files")[0]
    emb = db.q("SELECT COUNT(*) AS n FROM doc_chunks WHERE emb IS NOT NULL")[0]["n"]
    return {**r, "with_meaning": emb, "running": STATE["running"], "note": STATE["last_note"],
            "folders": ROOTS, "embed_model": embed_model()}


# ---------------------------------------------------------------- searching
def _keyword(question, k):
    words = [w for w in re.findall(r"\w+", question.lower()) if len(w) > 2][:12]
    if not words:
        return []
    try:
        return [r["chunk_id"] for r in db.q(
            "SELECT chunk_id FROM doc_fts WHERE doc_fts MATCH ? ORDER BY bm25(doc_fts) LIMIT ?",
            (" OR ".join(f'"{w}"' for w in words), k))]
    except Exception:
        like = [r["id"] for r in db.q("SELECT id FROM doc_chunks WHERE text LIKE ? LIMIT ?", (f"%{words[0]}%", k))]
        return like


def _semantic(question, k):
    qv = embed([question], "query")
    if not qv:
        return []
    stamp = tuple(db.q("SELECT COUNT(emb) AS n, MAX(id) AS m FROM doc_chunks")[0].values())
    if _cache.get("stamp") != stamp:
        _cache.update(stamp=stamp, rows=db.q("SELECT id, emb FROM doc_chunks WHERE emb IS NOT NULL"), m=None)
    rows = _cache["rows"]
    if not rows:
        return []
    try:
        import numpy as np
        if _cache["m"] is None:
            _cache["m"] = np.array([_unpack(r["emb"]) for r in rows], dtype="float32")
        m = _cache["m"]
        q = np.array(qv[0], dtype="float32")
        sims = m @ q / (np.linalg.norm(m, axis=1) * np.linalg.norm(q) + 1e-9)
        order = np.argsort(-sims)[:k]
        return [rows[i]["id"] for i in order if sims[i] > 0.35]
    except ImportError:
        q = qv[0]
        qn = sum(x * x for x in q) ** 0.5
        scored = []
        for r in rows:
            v = _unpack(r["emb"])
            s = sum(a * b for a, b in zip(v, q)) / ((sum(a * a for a in v) ** 0.5) * qn + 1e-9)
            scored.append((s, r["id"]))
        scored.sort(reverse=True)
        return [i for s, i in scored[:k] if s > 0.35]


def search(question, k=6):
    sem, kw = _semantic(question, 20), _keyword(question, 20)
    score = {}
    for lst in (sem, kw):  # reciprocal rank fusion: good in either list = near the top
        for rank, cid in enumerate(lst):
            score[cid] = score.get(cid, 0) + 1 / (60 + rank)
    best = sorted(score, key=score.get, reverse=True)[:k]
    if not best:
        return []
    rows = {r["id"]: r for r in db.q(f"SELECT id, path, chunk_no, text FROM doc_chunks WHERE id IN "
                                     f"({','.join('?' * len(best))})", best)}
    return [rows[i] for i in best if i in rows]


@tool("search_my_files",
      "Search INSIDE all the user's documents (Word, PDF, Excel, text files in their folders) and return the most "
      "relevant passages with file names. Use this whenever the user asks about something that may be written in "
      "their own files, e.g. 'what did the contract say about payment', 'find the invoice from Rahim'.",
      {"question": S("What to look for, in natural words."), "how_many": I("Number of passages (default 6).")},
      ["question"])
def search_my_files(question, how_many=6):
    if not db.q("SELECT 1 FROM doc_files LIMIT 1"):
        threading.Thread(target=index_files, daemon=True).start()
        return "Your files haven't been indexed yet. I started indexing now; try again in a few minutes."
    hits = search(question, max(1, min(int(how_many or 6), 12)))
    if not hits:
        return "Nothing relevant found in your indexed files."
    out = [f"[{n}] {h['path']} (part {h['chunk_no'] + 1})\n{h['text']}" for n, h in enumerate(hits, 1)]
    return _cut("Passages from your files (mention the file names in your answer):\n\n" + "\n\n".join(out), 10000)


@tool("index_my_files", "Re-scan the user's folders so search_my_files knows about new or changed documents.")
def index_my_files():
    return index_files()
