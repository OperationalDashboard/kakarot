"""Everything Kakarot can actually DO. Each @tool becomes a function the local model can call."""
import datetime
import difflib
import fnmatch
import json
import os
import re
import shutil
import uuid

import requests

import db
from config import CFG

TOOLS = {}
CTX = {"background": False}  # set by the agent for each run


def tool(name, description, params=None, required=None):
    params = params or {}

    def deco(fn):
        TOOLS[name] = {
            "fn": fn,
            "schema": {
                "type": "function",
                "function": {
                    "name": name,
                    "description": description,
                    "parameters": {
                        "type": "object",
                        "properties": params,
                        "required": list(params.keys()) if required is None else required,
                    },
                },
            },
        }
        return fn

    return deco


def S(d):
    return {"type": "string", "description": d}


def I(d):
    return {"type": "integer", "description": d}


def B(d):
    return {"type": "boolean", "description": d}


def schemas():
    return [t["schema"] for t in TOOLS.values()]


# ---------------------------------------------------------------- only the abilities a message needs
# The local brain's memory is small (8K tokens). Sending all 70+ ability descriptions every time would push the
# instructions and the chat out of its memory, so each message gets the always-needed ones plus the groups that
# match what the user is talking about (then others, while they fit).
CORE = ["list_files", "read_file", "write_file", "make_folder", "move_file", "trash_file", "search_files", "web_search",
        "fetch_page", "remember", "add_note", "set_reminder", "get_time", "start_mission"]
GROUPS = [  # (name, tools or module names, keywords)
    ("websites", ["sitebuilder", "workshop"],
     r"web ?site|\bsite\b|web ?app|web ?page|landing|portfolio|publish|deploy|github|\brepo|hosting|\bftp|cpanel|"
     r"namecheap|netlify|vercel|domain|\bseo\b|online|html|css|lalsobujdata|greenwonderland|"
     r"\.(com|site|net|org|io|app|dev|bd|xyz|shop)\b"),
    ("lalsobuj", ["lalsobuj_trip_report", "lalsobuj_route_trips", "lalsobuj_morning_brief"],
     r"\bbus|\btrips?\b|route|\bseats?\b|ticket|lal ?sobuj|dhaka|sylhet|chatt?og|chittagong|noakhali|counter|"
     r"coach|\bfares?\b|cox|rajshahi|khulna|barishal|comilla|feni"),
    ("code", ["coder"], r"\bcode|program|script|python|javascript|\bfunction|\bbug|\bapp\b|automat|\bbot\b|\bapi\b|"
                        r"test it|run it|debug|\.py\b|software"),
    ("office", ["office"], r"excel|xlsx|\bsheet|spreadsheet|\bword\b|docx|\bpdf|document|letter|\breport|invoice|"
                           r"\btable|chart|\bcsv"),
    ("browser", ["browser_ctl", "browse_page", "check_page_changes"],
     r"\bopen\b|click|log ?in|sign ?in|\bfill|\bform\b|button|browse|go to|type in|select|scroll|screenshot|"
     r"https?://|www\.|search on|book (a|the)|changed"),
    ("vision", ["vision"], r"photo|picture|image|screenshot|scan|look at|attached|\.(png|jpe?g|webp|gif)\b|see this"),
    ("files", ["knowledge", "file_info"], r"my files|my documents|in my|folder|search my|documents? say|agreement|"
                                          r"contract|find (the|my) file|index"),
    ("pc", ["pctools"], r"command|\bcmd\b|\brun\b|ip address|ipconfig|install|open (the |this |my )?(file|folder|app)|"
                        r"\bpc\b|computer|notepad|restart|shutdown"),
    ("watch", ["monitor"], r"watch|monitor|\bup\b|\bdown\b|uptime|brief|morning|certificate"),
    ("schedule", ["schedule_task", "list_scheduled", "cancel_scheduled"],
     r"remind|every (day|week|month|hour|morning|night)|schedul|daily|weekly|\bat \d|tomorrow|cancel"),
    ("memory", ["forget", "list_memory", "list_notes", "complete_note"], r"remember|forget|memory|\bnotes?\b|to-?do|done|tick"),
    ("missions", ["missions"], r"mission|background|research|compare|in detail|deep"),
    ("self", ["selfupdate", "plugins"], r"update (your)?self|new version|learn|new abilit|tool for|build (your)?self|"
                                        r"my tools|upgrade"),
]
FILL_ORDER = ["websites", "browser", "office", "code", "lalsobuj", "files", "vision", "missions", "pc", "schedule",
              "memory", "watch", "self"]


def _group_tools(spec):
    names = []
    for x in spec:
        if x in TOOLS:
            names.append(x)
        else:
            names += [n for n, t in TOOLS.items() if t["fn"].__module__ == x]
    return names


def schemas_for(messages, ctx_tokens=8192):
    """Ability descriptions for this conversation: core + matching groups, then more while they fit."""
    texts = [m.get("content") or "" for m in messages if m.get("role") == "user"][-3:]
    text = " ".join(texts).lower()
    used = {c["function"]["name"] for m in messages for c in (m.get("tool_calls") or []) if c.get("function")}
    budget = int(ctx_tokens * 4 * 0.42)
    chosen = [n for n in CORE if n in TOOLS] + [n for n in used if n in TOOLS and n not in CORE]
    size = lambda names: len(json.dumps([TOOLS[n]["schema"] for n in names]))
    groups = {g: _group_tools(spec) for g, spec, _ in GROUPS}
    matched = [g for g, _, kw in GROUPS if re.search(kw, text)]
    plugin_tools = [n for n, t in TOOLS.items() if t["fn"].__module__.startswith("plugins.") or
                    t["fn"].__module__.startswith("kakarot_plugin")]
    for g in matched + ["__plugins"] + [g for g in FILL_ORDER if g not in matched]:
        names = plugin_tools if g == "__plugins" else groups.get(g, [])
        extra = [n for n in names if n not in chosen]
        if not extra:
            continue
        if g in matched or g == "__plugins" or size(chosen + extra) <= budget:
            chosen += extra
    # anything else at all? (tools not in any group, e.g. new plugins) - add while it fits
    for n in TOOLS:
        if n not in chosen and size(chosen + [n]) <= budget:
            chosen.append(n)
    return [TOOLS[n]["schema"] for n in chosen]


def execute(name, args):
    t = TOOLS.get(name)
    if not t:
        return f"Error: unknown tool '{name}'. Available tools: {', '.join(TOOLS)}"
    try:
        out = t["fn"](**(args or {}))
        return out if isinstance(out, str) else json.dumps(out, ensure_ascii=False, default=str)
    except TypeError as e:
        return f"Error: wrong arguments for {name}: {e}"
    except Exception as e:  # report every failure back to the model so it can recover
        return f"Error: {type(e).__name__}: {e}"


# ---------------------------------------------------------------- file safety
ROOTS = CFG["allowed_folders"]
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"}


def _n(p):
    return os.path.normcase(os.path.realpath(p))


def safe_path(p):
    """Only allow paths inside the folders listed in config.json."""
    p = os.path.expandvars(os.path.expanduser(str(p).strip().strip('"')))
    if not os.path.isabs(p):
        p = os.path.join(ROOTS[0], p)
    rp = os.path.realpath(p)
    for r in ROOTS:
        if _n(rp) == _n(r) or _n(rp).startswith(_n(r) + os.sep):
            return rp
    raise PermissionError(f"'{p}' is outside the allowed folders {ROOTS}. "
                          f"Add it to allowed_folders in config.json if you need it.")


def _cut(text, limit):
    text = text or ""
    return text if len(text) <= limit else text[:limit] + f"\n...[cut, {len(text) - limit:,} more characters]"


# ---------------------------------------------------------------- files
@tool("list_files", "List the files and folders inside a folder.",
      {"path": S("Folder path. Relative paths start in the main workspace; '.' = workspace.")}, [])
def list_files(path="."):
    p = safe_path(path)
    rows = []
    for name in sorted(os.listdir(p))[:400]:
        fp = os.path.join(p, name)
        if os.path.isdir(fp):
            rows.append(f"[DIR]  {name}")
        else:
            rows.append(f"{name}  ({os.path.getsize(fp):,} bytes)")
    return f"{p}\n" + ("\n".join(rows) or "(empty folder)")


@tool("read_file", "Read a text, PDF or Word (.docx) file and return its text.",
      {"path": S("File path.")})
def read_file(path):
    p = safe_path(path)
    ext = os.path.splitext(p)[1].lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        text = "\n".join((pg.extract_text() or "") for pg in PdfReader(p).pages)
        if len(text.strip()) < 40:
            return "This PDF has no readable text (it's probably scanned). Use look_at_image on it to read it."
    elif ext == ".docx":
        import docx
        d = docx.Document(p)
        text = "\n".join(par.text for par in d.paragraphs)
        for t in d.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows)
    elif ext in (".xlsx", ".xlsm", ".csv"):
        import office
        return office.read_excel(path)
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"):
        return "This is a picture - use look_at_image to see it."
    else:
        with open(p, encoding="utf-8", errors="replace") as f:
            text = f.read()
    return _cut(text, 12000)


@tool("write_file", "Create or overwrite a text file (or append to it). Use .md or .txt for documents.",
      {"path": S("File path."), "content": S("Full text to write."),
       "append": B("true to add to the end instead of overwriting.")}, ["path", "content"])
def write_file(path, content, append=False):
    p = safe_path(path)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "a" if append else "w", encoding="utf-8") as f:
        f.write(content)
    return f"Saved {len(content):,} characters to {p}"


@tool("make_folder", "Create a folder (and parent folders).", {"path": S("Folder path.")})
def make_folder(path):
    p = safe_path(path)
    os.makedirs(p, exist_ok=True)
    return f"Folder ready: {p}"


@tool("move_file", "Move or rename a file or folder.",
      {"source": S("Current path."), "destination": S("New path (a folder or a full new file name).")})
def move_file(source, destination):
    src, dst = safe_path(source), safe_path(destination)
    if os.path.isdir(dst):
        dst = os.path.join(dst, os.path.basename(src))
    if os.path.exists(dst):
        raise FileExistsError(f"{dst} already exists")
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    shutil.move(src, dst)
    return f"Moved {src} -> {dst}"


@tool("trash_file", "Safely remove a file/folder by moving it to the workspace _trash folder (never deletes).",
      {"path": S("Path to remove.")})
def trash_file(path):
    src = safe_path(path)
    trash = os.path.join(ROOTS[0], "_trash")
    os.makedirs(trash, exist_ok=True)
    dst = os.path.join(trash, datetime.datetime.now().strftime("%Y%m%d-%H%M%S_") + os.path.basename(src))
    shutil.move(src, dst)
    return f"Moved to trash: {dst}"


@tool("search_files", "Find files by name pattern (like *.pdf or *invoice*), optionally containing some text.",
      {"pattern": S("File name pattern, e.g. *.pdf"), "path": S("Folder to search in ('.' = workspace)."),
       "contains": S("Optional text that must appear inside the file.")}, ["pattern"])
def search_files(pattern, path=".", contains=""):
    root = safe_path(path)
    hits = []
    for d, _, files in os.walk(root):
        for f in files:
            if not fnmatch.fnmatch(f.lower(), pattern.lower()):
                continue
            fp = os.path.join(d, f)
            if contains:
                try:
                    if os.path.getsize(fp) > 3_000_000:
                        continue
                    with open(fp, encoding="utf-8", errors="ignore") as fh:
                        if contains.lower() not in fh.read().lower():
                            continue
                except OSError:
                    continue
            hits.append(fp)
            if len(hits) >= 100:
                return "\n".join(hits) + "\n...(first 100 shown)"
    return "\n".join(hits) or "No matching files."


@tool("file_info", "Get size and modified date of a file.", {"path": S("File path.")})
def file_info(path):
    p = safe_path(path)
    st = os.stat(p)
    return {"path": p, "bytes": st.st_size,
            "modified": datetime.datetime.fromtimestamp(st.st_mtime).isoformat(sep=" ", timespec="seconds")}


# ---------------------------------------------------------------- web
@tool("web_search", "Search the internet (DuckDuckGo). Returns titles, links and snippets.",
      {"query": S("Search words."), "max_results": I("How many results (1-10).")}, ["query"])
def web_search(query, max_results=5):
    try:
        from ddgs import DDGS
    except ImportError:
        from duckduckgo_search import DDGS
    res = list(DDGS().text(query, max_results=max(1, min(int(max_results), 10))))
    if not res:
        return "No results."
    return "\n\n".join(f"{r.get('title')}\n{r.get('href')}\n{r.get('body')}" for r in res)


def _html_to_text(html):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "noscript", "svg"]):
        t.decompose()
    text = soup.get_text("\n")
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


@tool("fetch_page", "Quickly download a web page and return its readable text. Fast; use browse_page if the "
                    "page needs JavaScript or comes back empty.",
      {"url": S("Full URL starting with http.")})
def fetch_page(url):
    r = requests.get(url, headers=UA, timeout=30)
    r.raise_for_status()
    text = _html_to_text(r.text)
    if len(text) < 200:
        text += "\n\n(Very little text - this page probably needs JavaScript. Try browse_page.)"
    return _cut(text, 4000)


def _browse(url, screenshot=False, wait_seconds=3):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        try:
            page = browser.new_page(user_agent=UA["User-Agent"])
            page.goto(url, wait_until="domcontentloaded", timeout=45000)
            page.wait_for_timeout(int(float(wait_seconds) * 1000))
            title = page.title()
            text = page.inner_text("body")
            shot = ""
            if screenshot:
                folder = os.path.join(ROOTS[0], "screenshots")
                os.makedirs(folder, exist_ok=True)
                shot = os.path.join(folder, datetime.datetime.now().strftime("%Y%m%d-%H%M%S") + ".png")
                page.screenshot(path=shot, full_page=True)
            return title, text, shot
        finally:
            browser.close()


@tool("browse_page", "Open a web page in a real (hidden) Chrome browser, let JavaScript load, and return the "
                     "text. Can also save a screenshot into the workspace.",
      {"url": S("Full URL."), "screenshot": B("Save a full-page screenshot."),
       "wait_seconds": I("Seconds to wait for the page to load (default 3).")}, ["url"])
def browse_page(url, screenshot=False, wait_seconds=3):
    title, text, shot = _browse(url, screenshot, wait_seconds)
    out = f"Title: {title}\n\n{_cut(text, 4000)}"
    if shot:
        out += f"\n\nScreenshot saved: {shot}"
    return out


@tool("check_page_changes", "Check a web page and report what changed since the last time it was checked. "
                            "Perfect for monitoring prices, notices, stock, results, etc.",
      {"url": S("Full URL."), "use_browser": B("true for JavaScript-heavy sites.")}, ["url"])
def check_page_changes(url, use_browser=False):
    text = _browse(url)[1] if use_browser else _html_to_text(requests.get(url, headers=UA, timeout=30).text)
    old = db.q("SELECT text, checked FROM page_snapshots WHERE url=?", (url,))
    db.x("INSERT OR REPLACE INTO page_snapshots(url, text, checked) VALUES(?,?,?)", (url, text, db.now()))
    if not old:
        return f"First check saved (baseline). Page text now:\n{_cut(text, 3000)}"
    if old[0]["text"] == text:
        return f"No change since {old[0]['checked']}."
    diff = difflib.unified_diff(old[0]["text"].splitlines(), text.splitlines(), lineterm="", n=0)
    lines = [l for l in diff if l[:1] in "+-" and not l.startswith(("+++", "---"))]
    return f"CHANGED since {old[0]['checked']}:\n" + _cut("\n".join(lines), 4000)


@tool("lalsobuj_trip_report",
      "Check EVERY route/rotation on the Lal Sobuj bus website (lalsobujbus.com) and report which routes have "
      "trips, how many, and which have none, for one date. Does everything in one step (takes 1-4 minutes). "
      "Use this for any question about Lal Sobuj trips, rotations or routes.",
      {"date": S("'today', 'tomorrow' or YYYY-MM-DD (within the next week). Default today.")}, [])
def lalsobuj_trip_report(date="today"):
    import bus_report
    return "[[DIRECT]]" + bus_report.report(
        date, save_folder=os.path.join(ROOTS[0], "reports"))


@tool("lalsobuj_route_trips",
      "Find the trips for ONE Lal Sobuj bus route (origin -> destination) on a date, with bus code, departure "
      "time, counter, coach type, seats left and fare. Use when the user names an origin and a destination.",
      {"origin": S("From city, e.g. Dhaka"), "destination": S("To city, e.g. Noakhali"),
       "date": S("'today', 'tomorrow' or YYYY-MM-DD (within the next week). Default today.")},
      ["origin", "destination"])
def lalsobuj_route_trips(origin, destination, date="today"):
    import bus_report
    return "[[DIRECT]]" + bus_report.route(origin, destination, date)


# ---------------------------------------------------------------- memory & notes
@tool("remember", "Save a long-term fact about the user or their preferences (it will be remembered forever).",
      {"fact": S("The fact, e.g. 'User's office starts at 9am'.")})
def remember(fact):
    db.x("INSERT INTO memory(fact, created) VALUES(?,?)", (fact, db.now()))
    return "Remembered."


@tool("forget", "Delete a remembered fact by its id.", {"fact_id": I("Id from list_memory.")})
def forget(fact_id):
    db.x("DELETE FROM memory WHERE id=?", (int(fact_id),))
    return "Forgotten."


@tool("list_memory", "Show all remembered facts with ids.")
def list_memory():
    return db.q("SELECT id, fact FROM memory ORDER BY id") or "Nothing remembered yet."


@tool("add_note", "Add a note or to-do item.", {"text": S("Note text.")})
def add_note(text):
    nid = db.x("INSERT INTO notes(text, created) VALUES(?,?)", (text, db.now()))
    return f"Note #{nid} added."


@tool("list_notes", "List notes / to-dos.", {"include_done": B("Also show finished items.")}, [])
def list_notes(include_done=False):
    sql = "SELECT id, text, done, created FROM notes" + ("" if include_done else " WHERE done=0") + " ORDER BY id"
    return db.q(sql) or "No notes."


@tool("complete_note", "Mark a to-do as done.", {"note_id": I("Note id.")})
def complete_note(note_id):
    db.x("UPDATE notes SET done=1 WHERE id=?", (int(note_id),))
    return f"Note #{note_id} marked done."


# ---------------------------------------------------------------- reminders & schedules
def _parse_when(when):
    w = when.strip().lower()
    m = re.fullmatch(r"\+?\s*(\d+)\s*(m|min|mins|minutes?|h|hrs?|hours?|d|days?)", w)
    if m:
        n, u = int(m.group(1)), m.group(2)[0]
        delta = {"m": datetime.timedelta(minutes=n), "h": datetime.timedelta(hours=n),
                 "d": datetime.timedelta(days=n)}[u]
        return datetime.datetime.now() + delta
    return datetime.datetime.fromisoformat(when.strip().replace("T", " "))


@tool("set_reminder", "Set a one-time reminder. It pops up on the PC and in the Inbox at that time.",
      {"message": S("What to remind about."),
       "when": S("Local date-time 'YYYY-MM-DD HH:MM', or relative like '30m', '2h', '1d'.")})
def set_reminder(message, when):
    import scheduler
    at = _parse_when(when)
    if at < datetime.datetime.now():
        raise ValueError(f"{at} is in the past")
    jid = "r-" + uuid.uuid4().hex[:8]
    scheduler.add_job({"id": jid, "name": message, "prompt": message, "cron": None,
                       "run_at": at.isoformat(sep=" ", timespec="minutes"), "kind": "reminder"})
    return f"Reminder {jid} set for {at:%a %d %b %Y %H:%M}."


@tool("schedule_task", "Make Kakarot do a task automatically on a repeating schedule (e.g. check a website every "
                       "morning). Results go to the Inbox.",
      {"name": S("Short name."),
       "task": S("Full instruction Kakarot should carry out each time, written as if the user typed it."),
       "cron": S("5-field cron in local time: 'minute hour day month weekday'. "
                 "Examples: '0 9 * * *' daily 9:00, '*/30 * * * *' every 30 min, '0 8 * * 0-4' Sun-Thu 8:00.")})
def schedule_task(name, task, cron):
    import scheduler
    jid = "t-" + uuid.uuid4().hex[:8]
    nxt = scheduler.add_job({"id": jid, "name": name, "prompt": task, "cron": cron.strip(),
                             "run_at": None, "kind": "task"})
    return f"Scheduled '{name}' ({jid}). Next run: {nxt}."


@tool("list_scheduled", "List upcoming reminders and scheduled tasks.")
def list_scheduled():
    import scheduler
    return scheduler.list_jobs() or "Nothing scheduled."


@tool("cancel_scheduled", "Cancel a reminder or scheduled task by id.", {"job_id": S("Id like t-1a2b3c4d.")})
def cancel_scheduled(job_id):
    import scheduler
    scheduler.remove_job(job_id)
    return f"Cancelled {job_id}."


@tool("get_time", "Get the current local date and time.")
def get_time():
    return datetime.datetime.now().strftime("%A %d %B %Y, %H:%M")


# ---------------------------------------------------------------- extra abilities (each file adds its own tools)
import browser_ctl  # noqa: E402,F401  website control: click / type / fill forms
import office  # noqa: E402,F401       Excel / Word / PDF
import pctools  # noqa: E402,F401      run commands on the PC (with approval)
import vision  # noqa: E402,F401      see pictures and scanned PDFs
import knowledge  # noqa: E402,F401   search inside the user's documents
import missions  # noqa: E402,F401    big background jobs
import monitor  # noqa: E402,F401     morning brief + website watchdog
import coder  # noqa: E402,F401      write / run / fix real programs
import selfupdate  # noqa: E402,F401 nightly learning, new versions, brain refresh
import sitebuilder  # noqa: E402,F401 complete websites: plan, write, browser-test, fix, preview
import workshop  # noqa: E402,F401   open any website, deep check, fix/change safely, publish (GitHub/FTP/Netlify/Vercel)
import plugins  # noqa: E402         abilities Kakarot built for itself (approved by the user)
plugins.load_all()
